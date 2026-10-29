import sys, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
src, out = sys.argv[1], sys.argv[2]
df = pd.read_csv(src, thousands=','); df.columns = df.columns.str.strip()
df['MY'] = df.ReportingMarketingYearName.str.extract(r'(\d{4})/Sep (\d{4})').agg('/'.join, axis=1)
df['Date'] = pd.to_datetime(df.WeekendingDate)
df = df[df.Date >= pd.to_datetime(df.MY.str[:4] + '-10-01')]   # drop forward-year duplicate rows
df = df[df.MY != '2019/2020']                                    # only its straddle week is present
wk = df.groupby(['MY', 'WeekNumber']).agg(ns=('TotalNetSales', 'first'), tc=('TotalTotalCommitment', 'first'))
mys = sorted(wk.index.levels[0].intersection(df.MY.unique()))
F = 'Arial'; bold = Font(name=F, bold=True); norm = Font(name=F); white = Font(name=F, bold=True, color='FFFFFF')
hdr = PatternFill('solid', fgColor='1F4E78'); NUM = '#,##0;(#,##0);-'
wb = Workbook(); wb.remove(wb.active)
def sheet(name, title, note):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(name=F, bold=True, size=12)
    ws['A2'] = note; ws['A2'].font = Font(name=F, italic=True, size=9)
    for j, h in enumerate(['Week'] + mys, 1):
        c = ws.cell(3, j, h); c.font = white; c.fill = hdr; c.alignment = Alignment(horizontal='center')
        ws.column_dimensions[c.column_letter].width = 9 if j == 1 else 13
    for w in range(1, 54): ws.cell(3 + w, 1, w).font = bold
    ws.freeze_panes = 'B4'
    return ws
note_src = 'Metric tons. Source: USDA FAS Export Sales, European destinations (user download). Marketing year Oct–Sep. Blank = week not yet reported.'
w = sheet('Weekly net sales', 'Total EU — weekly net sales, current MY (t)', note_src)
c = sheet('Cumulative net sales', 'Total EU — cumulative net sales since week 1 (t)', 'Running sum of "Weekly net sales". Excludes sales booked before Oct 1 for this MY — see "Total commitments".')
t = sheet('Total commitments', 'Total EU — total commitments (accumulated exports + outstanding sales) (t)', 'USDA reported figure. = forward sales booked before the MY started + cumulative net sales. Standard measure of sales progression.')
for j, my in enumerate(mys, 2):
    col = w.cell(3, j).column_letter
    for wn in range(1, 54):
        r = 3 + wn
        if (my, wn) in wk.index:
            w.cell(r, j, float(wk.loc[(my, wn), 'ns'])); t.cell(r, j, float(wk.loc[(my, wn), 'tc']))
        prev = '' if wn == 1 else f'N({col}{r-1})+'
        c.cell(r, j, f"=IF('Weekly net sales'!{col}{r}=\"\",\"\",{prev}'Weekly net sales'!{col}{r})")
        for ws in (w, c, t): ws.cell(r, j).number_format = NUM; ws.cell(r, j).font = norm
wb.save(out)
