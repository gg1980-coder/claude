import sys, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

src, out = sys.argv[1], sys.argv[2]
df = pd.read_csv(src, thousands=',')
df.columns = df.columns.str.strip()
df['Country'] = df.CountryName.str.strip()
df['MY'] = df.ReportingMarketingYearName.str.replace('MY -', '').str.strip().str.replace(r'Oct (\d{4})/Sep (\d{4})', r'\1/\2', regex=True)
df['Date'] = pd.to_datetime(df.WeekendingDate)
df['RowType'] = 'Current MY'
df.loc[df.Date < pd.to_datetime(df.MY.str[:4] + '-10-01'), 'RowType'] = 'Forward (dup)'
# straddle week (w/e on/after Oct 1) is also reported as week 1 of the next MY -> don't count its NextMY sales twice
nx = df.NextMYNetSales.where(df.Date < pd.to_datetime(df.MY.str[5:] + '-10-01'), 0)
df['NetSales_incl_NextMY'] = (df.NetSales + nx).where(df.RowType == 'Current MY', 0)
metrics = ['NetSales', 'NextMYNetSales', 'NetSales_incl_NextMY', 'GrossNewSales', 'WeeklyExports']
data = df[['MY', 'WeekNumber', 'Date', 'Country', 'RowType'] + metrics].sort_values(['Date', 'MY', 'Country']).reset_index(drop=True)
mys = sorted(m for m in data.MY.unique() if m != '2019/2020')  # 2019/20 only has its straddle week 53
countries = sorted(data.Country.unique())

F = 'Arial'
bold = Font(name=F, bold=True); norm = Font(name=F); white = Font(name=F, bold=True, color='FFFFFF')
hdr = PatternFill('solid', fgColor='1F4E78'); tot = PatternFill('solid', fgColor='DDEBF7'); inp = PatternFill('solid', fgColor='FFFF00')
thin = Border(top=Side(style='thin'))
NUM = '#,##0;(#,##0);-'

wb = Workbook()
rd = wb.active; rd.title = 'README'
n = len(data) + 1
# ---- Data sheet
ds = wb.create_sheet('Data')
cols = ['MY', 'Week', 'WeekEnding', 'Country', 'RowType'] + metrics
ds.append(cols)
for r in data.itertuples(index=False):
    ds.append([r.MY, r.WeekNumber, r.Date.to_pydatetime(), r.Country, r.RowType] + [getattr(r, m) for m in metrics])
for c in ds[1]: c.font = white; c.fill = hdr
for row in ds.iter_rows(min_row=2):
    for c in row: c.font = norm
    row[2].number_format = 'yyyy-mm-dd'
    for c in row[5:]: c.number_format = NUM
for i, w in enumerate([11, 7, 12, 22, 15] + [20] * len(metrics), 1): ds.column_dimensions[L(i)].width = w
ds.freeze_panes = 'A2'; ds.auto_filter.ref = f'A1:{L(len(cols))}{n}'
lastm = L(5 + len(metrics))
rngMY, rngW, rngC = f'Data!$A$2:$A${n}', f'Data!$B$2:$B${n}', f'Data!$D$2:$D${n}'
rngT = f'Data!$E$2:$E${n}'
rngVal = f'Data!$F$2:${lastm}${n}'
metricCell = 'README!$B$4'
valCol = f'INDEX({rngVal},0,MATCH({metricCell},Data!$F$1:${lastm}$1,0))'

# ---- README
rd['A1'] = 'US soybean cake & meal export sales to Europe — weekly, by marketing year'; rd['A1'].font = Font(name=F, bold=True, size=13)
rd['A3'] = 'Choose the metric shown on every country/total sheet (yellow cell, dropdown):'; rd['A3'].font = bold
rd['A4'] = 'Metric'; rd['A4'].font = bold
rd['B4'] = 'NetSales'; rd['B4'].fill = inp; rd['B4'].font = bold
dv = DataValidation(type='list', formula1='"' + ','.join(metrics) + '"', allow_blank=False); rd.add_data_validation(dv); dv.add('B4')
notes = [
 ('NetSales', 'Net new sales for the CURRENT marketing year booked in the week (gross sales minus cancellations/reductions). USDA headline "net sales".'),
 ('NextMYNetSales', 'Net sales booked in the week for delivery in the NEXT marketing year (forward sales, mostly Jun–Sep).'),
 ('NetSales_incl_NextMY', 'NetSales + NextMYNetSales = all new business booked that week, regardless of delivery year. Straddle week (w/e on/after 1 Oct, reported in both MYs) counted once.'),
 ('GrossNewSales', 'New sales before cancellations.'),
 ('WeeklyExports', 'Physical shipments in the week (not sales).'),
 ('', ''),
 ('Units', 'Metric tons.'),
 ('Year', 'USDA marketing year Oct–Sep (e.g. 2025/2026 = Oct 2025–Sep 2026). Week 1 = first week of October; some years have 53 weeks.'),
 ('Blanks', 'Blank cell = week not yet reported. Zero/"-" = reported week with no activity for that country.'),
 ('Excluded', 'MY 2019/2020 appears in the source only as its final straddle week (w/e 2020-10-01); it is kept in Data but not shown in the grids.'),
 ('Forward rows', 'The USDA export also lists next-MY sales as rows labelled with the NEXT marketing year but dated in the current one (e.g. "MY 2026/27, week 41" dated 9 Jul 2026). Their week numbers collide with the real weeks a year later, so a naive pivot double-counts them. They are flagged RowType = "Forward (dup)" in Data and excluded from the grids; the same tonnage is already in NextMYNetSales.'),
 ('Check', 'Sum of countries reconciles exactly with USDA "Total" columns for every week in the source.'),
 ('Source', 'USDA FAS Export Sales Reporting query export (ESRQueryParameterReport), European destinations only, downloaded by user.'),
]
for i, (a, b) in enumerate(notes, 6):
    rd.cell(i, 1, a).font = bold; rd.cell(i, 2, b).font = norm
rd.column_dimensions['A'].width = 24; rd.column_dimensions['B'].width = 120
rd['A20'] = 'Sheets'; rd['A20'].font = bold
rd['B20'] = 'TOTAL = all European destinations; one sheet per country; Data = clean long table (filterable / pivotable); Week dates = week-ending date for each week & year.'; rd['B20'].font = norm

def grid(ws, country):
    ws['A1'] = f'{country} — '; ws['A1'].font = Font(name=F, bold=True, size=12)
    ws['A1'] = f'="{country} — "&{metricCell}&" (metric tons)"'
    ws.cell(3, 1, 'Week').font = white; ws.cell(3, 1).fill = hdr
    for j, my in enumerate(mys, 2):
        c = ws.cell(3, j, my); c.font = white; c.fill = hdr; c.alignment = Alignment(horizontal='center')
        ws.column_dimensions[L(j)].width = 13
    ws.column_dimensions['A'].width = 9
    for w in range(1, 54):
        r = 3 + w
        ws.cell(r, 1, w).font = bold
        for j, my in enumerate(mys, 2):
            col = L(j)
            reported = f'COUNTIFS({rngMY},{col}$3,{rngW},$A{r},{rngT},"Current MY")=0'
            crit = f'{rngMY},{col}$3,{rngW},$A{r},{rngT},"Current MY"' + ('' if country == 'TOTAL' else f',{rngC},"{country}"')
            c = ws.cell(r, j, f'=IF({reported},"",SUMIFS({valCol},{crit}))')
            c.number_format = NUM; c.font = norm
    tr = 57
    ws.cell(tr, 1, 'Total').font = bold
    for j in range(1, len(mys) + 2):
        c = ws.cell(tr, j); c.fill = tot; c.border = thin
        if j > 1:
            c.value = f'=SUM({L(j)}4:{L(j)}56)'; c.font = bold; c.number_format = NUM
    ws.freeze_panes = 'B4'

grid(wb.create_sheet('TOTAL'), 'TOTAL')
for ctry in countries:
    grid(wb.create_sheet(ctry[:31]), ctry)

# ---- week dates reference
wd = wb.create_sheet('Week dates')
wd.cell(1, 1, 'Week').font = white; wd.cell(1, 1).fill = hdr
dates = data.groupby(['MY', 'WeekNumber']).Date.first()
for j, my in enumerate(mys, 2):
    c = wd.cell(1, j, my); c.font = white; c.fill = hdr; wd.column_dimensions[L(j)].width = 13
    for w in range(1, 54):
        if (my, w) in dates.index:
            c = wd.cell(1 + w, j, dates[(my, w)].to_pydatetime()); c.number_format = 'dd-mmm-yy'; c.font = norm
for w in range(1, 54): wd.cell(1 + w, 1, w).font = bold
wd.freeze_panes = 'B2'
wb.move_sheet('Data', offset=len(wb.sheetnames))
wb.save(out)
print(mys, countries)
